import base64
import json
import logging
import os
import re
from typing import Optional

import anthropic
from dotenv import load_dotenv
from pydantic import BaseModel

load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))

log = logging.getLogger(__name__)

ALLOWED_PDF = "application/pdf"
ALLOWED_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_PARSER_MODEL = "claude-haiku-4-5"

_RESUME_SCHEMA = """{
  "name": "string",
  "email": "string or null",
  "phone": "string or null",
  "location": "string or null",
  "links": { "linkedin": "...", "github": "...", "portfolio": "...", "other": ["..."] },
  "current_status": "string (e.g. 'Senior ML engineer at Acme')",
  "school_year": "string or null",
  "education": [{ "school": "...", "degree": "...", "year": "...", "gpa": "..." }],
  "work_experience": [{ "title": "...", "company": "...", "duration": "...", "bullets": ["..."], "metrics": ["..."] }],
  "skills": ["..."],
  "projects": [{ "name": "...", "description": "...", "metrics": ["..."] }],
  "publications": [{ "title": "...", "venue": "...", "year": "..." }],
  "research_interests": ["..."],
  "summary_one_line": "string"
}"""

PARSER_PROMPT = (
    "You are a resume parser. Extract structured data from the resume below.\n"
    "Return ONLY valid JSON matching this schema. No prose, no markdown.\n\n"
    + _RESUME_SCHEMA + "\n\n"
    "Rules:\n"
    "- Be precise. Quote bullets verbatim.\n"
    "- Pull metrics as separate from bullets so the writer agent can use them directly.\n"
    "- For \"summary_one_line\", write one sentence that would fit at the top of a cold email.\n"
    "- If a field has no data in the resume, return null or empty array, never invent.\n"
    "- Return ONLY the JSON object. No code fences, no commentary."
)


# ── Pydantic models ────────────────────────────────────────────────────────────

class ResumeLinks(BaseModel):
    linkedin: Optional[str] = None
    github: Optional[str] = None
    portfolio: Optional[str] = None
    other: list[str] = []


class ResumeEducation(BaseModel):
    school: Optional[str] = None
    degree: Optional[str] = None
    year: Optional[str] = None
    gpa: Optional[str] = None


class ResumeWorkExperience(BaseModel):
    title: Optional[str] = None
    company: Optional[str] = None
    duration: Optional[str] = None
    bullets: list[str] = []
    metrics: list[str] = []


class ResumeProject(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    metrics: list[str] = []


class ResumePublication(BaseModel):
    title: Optional[str] = None
    venue: Optional[str] = None
    year: Optional[str] = None


class ParsedResume(BaseModel):
    name: str
    email: Optional[str] = None
    phone: Optional[str] = None
    location: Optional[str] = None
    links: ResumeLinks = ResumeLinks()
    current_status: Optional[str] = None
    school_year: Optional[str] = None
    education: list[ResumeEducation] = []
    work_experience: list[ResumeWorkExperience] = []
    skills: list[str] = []
    projects: list[ResumeProject] = []
    publications: list[ResumePublication] = []
    research_interests: list[str] = []
    summary_one_line: str


# ── Core parser function ───────────────────────────────────────────────────────

def parse_resume(file_bytes: bytes, mime_type: str) -> dict:
    """
    Parse a PDF or DOCX resume and return structured JSON via Claude Haiku.

    Args:
        file_bytes: Raw bytes of the uploaded file.
        mime_type:  MIME type, must be ALLOWED_PDF or ALLOWED_DOCX.

    Returns:
        Validated resume dict matching the ParsedResume schema.

    Raises:
        RuntimeError: If ANTHROPIC_API_KEY is missing.
        ValueError:   If the file type is unsupported, Claude returns invalid JSON,
                      or the response fails schema validation.
    """
    api_key = os.getenv("ANTHROPIC_API_KEY")
    if not api_key:
        raise RuntimeError("ANTHROPIC_API_KEY is not set")

    client = anthropic.Anthropic(api_key=api_key)

    if mime_type == ALLOWED_PDF:
        content = _build_pdf_content(file_bytes)
    elif mime_type == ALLOWED_DOCX:
        content = _build_docx_content(file_bytes)
    else:
        raise ValueError(
            f"Unsupported file type: {mime_type!r}. "
            f"Must be {ALLOWED_PDF!r} or {ALLOWED_DOCX!r}."
        )

    log.debug("Calling Claude %s for resume parsing", _PARSER_MODEL)
    response = client.messages.create(
        model=_PARSER_MODEL,
        max_tokens=2000,
        messages=[{"role": "user", "content": content}],
    )
    raw_text = response.content[0].text.strip()

    # Strip accidental code fences
    raw_text = re.sub(r"^```(?:json)?\s*\n?", "", raw_text)
    raw_text = re.sub(r"\n?```\s*$", "", raw_text)
    raw_text = raw_text.strip()

    try:
        data = json.loads(raw_text)
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"Claude returned invalid JSON for resume parse: {exc}. "
            f"Raw response (first 300 chars): {raw_text[:300]}"
        ) from exc

    try:
        parsed = ParsedResume.model_validate(data)
    except Exception as exc:
        raise ValueError(
            f"Resume parse response failed schema validation: {exc}"
        ) from exc

    return parsed.model_dump()


# ── Content builders ───────────────────────────────────────────────────────────

def _build_pdf_content(file_bytes: bytes) -> list:
    """Build Claude message content for a PDF file using native document support."""
    encoded = base64.standard_b64encode(file_bytes).decode("utf-8")
    return [
        {
            "type": "document",
            "source": {
                "type": "base64",
                "media_type": ALLOWED_PDF,
                "data": encoded,
            },
        },
        {
            "type": "text",
            "text": PARSER_PROMPT,
        },
    ]


def _build_docx_content(file_bytes: bytes) -> list:
    """Extract text from a DOCX file and build plain-text message content."""
    try:
        import io
        import docx  # python-docx
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is required to parse DOCX files. "
            "Install it with: pip install python-docx"
        ) from exc

    doc = docx.Document(io.BytesIO(file_bytes))
    parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in parts:
                    parts.append(text)

    resume_text = "\n".join(parts)
    return [
        {
            "type": "text",
            "text": f"{PARSER_PROMPT}\n\nRESUME TEXT:\n{resume_text}",
        }
    ]
